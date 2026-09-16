"""tablex_formula_graph handler."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from ...domain.formula import (
    build_dependency_graph,
    extract_formulas,
    find_circular,
)
from ..schemas import ToolCall, ToolResult
from ._base import HandlerSpec, _fail, _ok


def handle_formula_graph(tool_call: ToolCall, session: Any) -> ToolResult:
    file_id = tool_call.input["file_id"]
    sheet = tool_call.input["sheet"]
    cell = tool_call.input.get("cell")

    if file_id not in session.files:
        return _fail(f"文件不存在: {file_id}")
    src_path = Path(session.files[file_id]["path"])
    if not src_path.exists():
        return _fail(f"源文件已不存在: {src_path}")
    if src_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return _fail("formula_graph 仅支持 xlsx/xlsm")

    from openpyxl import load_workbook

    wb = load_workbook(str(src_path), data_only=False, read_only=True)
    try:
        if sheet not in wb.sheetnames:
            return _fail(f"工作表不存在: {sheet}")
        formulas = extract_formulas(wb, sheet)
        graph = build_dependency_graph(formulas)
        circulars = find_circular(graph)

        if cell:
            upstream = _upstream_of(graph, cell)
            downstream = _downstream_of(graph, cell)
            return _ok(
                f"{cell} 依赖 {len(upstream)} 个上游单元格，影响 {len(downstream)} 个下游",
                data={
                    "cell": cell,
                    "upstream": sorted(upstream),
                    "downstream": sorted(downstream),
                    "formula_count": len(formulas),
                    "circular_count": len(circulars),
                },
            )

        return _ok(
            f"提取 {len(formulas)} 个公式，{len(circulars)} 个循环依赖",
            data={
                "formulas": [{"cell": f.cell, "formula": f.formula} for f in formulas],
                "graph": graph,
                "circulars": circulars,
                "formula_count": len(formulas),
                "circular_count": len(circulars),
            },
        )
    finally:
        wb.close()


def _upstream_of(graph: dict[str, list[str]], target: str) -> set[str]:
    """Cells that `target` (transitively) depends on."""
    seen: set[str] = set()
    stack = [target]
    while stack:
        node = stack.pop()
        for dep in graph.get(node, []):
            if dep in seen:
                continue
            seen.add(dep)
            stack.append(dep)
    return seen


def _downstream_of(graph: dict[str, list[str]], target: str) -> set[str]:
    """Cells that (transitively) depend on `target`."""
    reverse: dict[str, list[str]] = {}
    for src, deps in graph.items():
        for d in deps:
            reverse.setdefault(d, []).append(src)
    seen: set[str] = set()
    stack = [target]
    while stack:
        node = stack.pop()
        for src in reverse.get(node, []):
            if src in seen:
                continue
            seen.add(src)
            stack.append(src)
    return seen


HANDLERS = {
    "tablex_formula_graph": HandlerSpec(
        "tablex_formula_graph", ["file_id", "sheet"], handle_formula_graph,
    ),
}
