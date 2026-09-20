"""tablex_template_fill handler."""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd

from ...db import OutputRecord, init_db, insert_output
from ...domain.template import TemplateError, fill_template
from ..schemas import ToolCall, ToolResult
from ._base import HandlerSpec, _db_path, _fail, _ok, _output_id


def handle_template_fill(tool_call: ToolCall, session: Any) -> ToolResult:
    template_file_id = tool_call.input["template_file_id"]
    data = tool_call.input.get("data") or {}
    sheet = tool_call.input.get("sheet")

    if not isinstance(data, dict):
        return _fail("data 必须是 dict")
    if template_file_id not in session.files:
        return _fail(f"模板文件不存在: {template_file_id}")
    src_path = Path(session.files[template_file_id]["path"])
    if not src_path.exists():
        return _fail(f"模板文件已不存在: {src_path}")
    if src_path.suffix.lower() != ".xlsx":
        return _fail("template_fill 仅支持 .xlsx 模板")

    from openpyxl import load_workbook

    wb = load_workbook(str(src_path))
    try:
        sheet_name = sheet or (wb.sheetnames[0] if wb.sheetnames else None)
        if not sheet_name:
            return _fail("模板文件没有可用工作表")
        try:
            replaced, missing = fill_template(wb, sheet_name, data)
        except TemplateError as exc:
            return _fail(str(exc))

        output_id = _output_id(session)
        output_path = session.output_dir / f"{output_id}.xlsx"
        output_path.parent.mkdir(parents=True, exist_ok=True)
        # Persist the styled filled template as the primary artifact.
        wb.save(output_path)

        # Re-inject into session.tables so downstream tools can chain.
        reload = load_workbook(str(output_path), data_only=True, read_only=True)
        try:
            for sn in reload.sheetnames:
                rows = [list(r) for r in reload[sn].iter_rows(values_only=True)]
                if rows:
                    session.tables[f"{output_id}::{sn}"] = pd.DataFrame(rows)
        finally:
            reload.close()

        session.output_id = output_id
        session.output_path = str(output_path)

        db_path = _db_path()
        init_db(db_path)
        insert_output(
            db_path,
            OutputRecord(
                output_id=output_id,
                stored_path=str(output_path),
                source_file_ids=[template_file_id],
                plan={"tool": "tablex_template_fill", "input": tool_call.input},
                result={"replaced_count": replaced, "missing_keys": missing},
                status="completed",
                created_at=datetime.now(timezone.utc).isoformat(),
            ),
        )

        summary = f"模板填充完成：替换 {replaced} 处"
        if missing:
            summary += f"，缺失 {len(missing)} 个键：{missing[:5]}"
        return _ok(
            summary,
            data={
                "output_id": output_id,
                "replaced_count": replaced,
                "missing_keys": missing,
            },
        )
    finally:
        wb.close()


HANDLERS = {
    "tablex_template_fill": HandlerSpec(
        "tablex_template_fill", ["template_file_id", "data"], handle_template_fill,
    ),
}
