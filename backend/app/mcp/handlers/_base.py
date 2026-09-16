"""Shared helpers for tablex tool handlers.

Each per-tool module imports from here; handlers MUST NOT import from one another.
"""
from __future__ import annotations

import json
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook

from ...config import get_settings
from ...db import OutputRecord, init_db, insert_output
from ...domain.audit import make_audit_event
from ...domain.executor import INTERNAL_COL
from ..schemas import MAX_RESULT_BYTES, ToolCall, ToolResult, ValidationResult

# ----- Registry -----

HandlerFn = Callable[[ToolCall, "Any"], ToolResult]


class HandlerSpec:
    """Declares what fields a tool requires and which function runs it."""

    def __init__(self, name: str, required: list[str], fn: HandlerFn):
        self.name = name
        self.required = required
        self.fn = fn


# Tools that take file_id + sheet (sheet must already be loaded).
_FILE_SHEET_TOOLS = {
    "tablex_normalize",
    "tablex_deduplicate",
    "tablex_filter",
    "tablex_group_summary",
    "tablex_fill_formula",
    "tablex_sort",
    "tablex_fill_null",
    "tablex_pivot",
    "tablex_validate",
    "tablex_chart",
}

# Tools that take file_id only.
_FILE_TOOLS = {"tablex_inspect"}


# registry is set by __init__.py after auto-scan; each handler references the
# local HANDLERS dict at module load time, but validate_tool_input needs the
# global registry to look up required fields. To keep the function identical
# to the legacy single-file implementation, we accept HANDLERS via a setter
# instead of importing from __init__ (which would create a cycle).
_HANDLERS_REF: dict[str, HandlerSpec] | None = None


def set_handlers_registry(handlers: dict[str, HandlerSpec]) -> None:
    """Called once from handlers/__init__.py after auto-scan."""
    global _HANDLERS_REF
    _HANDLERS_REF = handlers


def get_handlers_registry() -> dict[str, HandlerSpec]:
    if _HANDLERS_REF is None:
        raise RuntimeError("Handler registry not initialized")
    return _HANDLERS_REF


def validate_tool_input(name: str, input_dict: dict, session: Any) -> ValidationResult:
    """Anti-corruption layer 1: parameter validation. Never enters the domain layer."""
    handlers = get_handlers_registry()
    if name not in handlers:
        return ValidationResult(ok=False, error=f"未知工具: {name}")

    spec = handlers[name]
    for field in spec.required:
        if field not in input_dict or input_dict[field] in (None, ""):
            return ValidationResult(ok=False, error=f"缺少参数: {field}")

    if name in _FILE_SHEET_TOOLS or name in _FILE_TOOLS:
        file_id = input_dict.get("file_id")
        if not session.has_file(file_id):
            return ValidationResult(ok=False, error=f"文件不存在: {file_id}")

    if name in _FILE_SHEET_TOOLS:
        sheet = input_dict.get("sheet")
        if sheet is None:
            return ValidationResult(ok=False, error="缺少参数: sheet")
        sheet_ref = f"{file_id}::{sheet}"
        if sheet_ref not in session.tables:
            return ValidationResult(ok=False, error=f"工作表未加载: {sheet}")

    if name == "tablex_compare":
        for ref_field in ("left_ref", "right_ref"):
            ref = input_dict.get(ref_field)
            if ref not in session.tables:
                return ValidationResult(ok=False, error=f"工作表未加载: {ref}")
        for col in input_dict.get("key_columns") or []:
            left_ref = input_dict["left_ref"]
            if col not in session.tables[left_ref].columns:
                return ValidationResult(ok=False, error=f"列 {col} 不在 {left_ref}")
            right_ref = input_dict["right_ref"]
            if col not in session.tables[right_ref].columns:
                return ValidationResult(ok=False, error=f"列 {col} 不在 {right_ref}")

    if name == "tablex_join":
        for side in ("left", "right"):
            ref = input_dict.get(side)
            if not isinstance(ref, dict):
                return ValidationResult(ok=False, error=f"{side} 必须是 {{file_id, sheet}}")
            file_id = ref.get("file_id")
            sheet = ref.get("sheet")
            if not file_id or not sheet:
                return ValidationResult(ok=False, error=f"{side} 缺少 file_id 或 sheet")
            if not session.has_file(file_id):
                return ValidationResult(ok=False, error=f"文件不存在: {file_id}")
            sheet_ref = f"{file_id}::{sheet}"
            if sheet_ref not in session.tables:
                return ValidationResult(ok=False, error=f"工作表未加载: {sheet_ref}")

    return ValidationResult(ok=True)


def truncate_result(result: ToolResult, max_bytes: int = MAX_RESULT_BYTES) -> str:
    """Serialize ToolResult; if over budget, drop the data field (keep summary/error)."""
    text = result.model_dump_json(exclude_none=True)
    if len(text.encode("utf-8")) <= max_bytes:
        return text
    minimal = {k: v for k, v in {"success": result.success, "summary": result.summary, "error": result.error}.items() if v is not None}
    return json.dumps(minimal, ensure_ascii=False)


def _ok(summary: str, data: dict[str, Any] | None = None) -> ToolResult:
    return ToolResult(success=True, summary=summary, data=data)


def _fail(error: str, summary: str = "处理失败") -> ToolResult:
    return ToolResult(success=False, summary=summary, error=error[:300])


def _audit(
    session: Any,
    operation: str,
    sheet_ref: str,
    columns: list[str],
    input_rows: int,
    output_rows: int,
    affected: list[int],
    details: dict[str, Any] | None = None,
) -> None:
    event = make_audit_event(
        step_id=f"step-{len(session.audit_events) + 1:03d}",
        operation=operation,
        input_sheets=[sheet_ref],
        output_sheets=[sheet_ref],
        columns=columns,
        input_rows=input_rows,
        output_rows=output_rows,
        affected_rows=affected,
        details=details,
    )
    session.audit_events.append(event)


def _db_path() -> Path:
    base = Path(get_settings().APP_DATA_DIR).resolve()
    base.mkdir(parents=True, exist_ok=True)
    return base / "metadata.db"


def _write_export_workbook(
    output_path: Path,
    tables: dict,
    audit_events: list,
) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    taken: set[str] = set()

    def _escape(name: str) -> str:
        # Openpyxl forbids these chars in sheet titles: :\\/?*[] and 31-char limit.
        cleaned = (
            name.replace("::", "_of_")
            .replace(":", "_")
            .replace("\\", "_")
            .replace("/", "_")
            .replace("?", "_")
            .replace("*", "_")
            .replace("[", "_")
            .replace("]", "_")
            .strip()
        ) or "工作表"
        candidate = cleaned[:31]
        suffix = 1
        while candidate in taken:
            suffix += 1
            stem = cleaned[: max(0, 31 - len(str(suffix)) - 1)]
            candidate = f"{stem}_{suffix}"
        taken.add(candidate)
        return candidate

    for ref, df in tables.items():
        ws = wb.create_sheet(_escape(ref))
        cleaned = df.drop(columns=[INTERNAL_COL], errors="ignore") if INTERNAL_COL in df.columns else df
        for col_idx, name in enumerate(cleaned.columns, start=1):
            ws.cell(row=1, column=col_idx, value=str(name)).font = ws.cell(row=1, column=col_idx).font.copy(bold=True)
        for row in cleaned.itertuples(index=False, name=None):
            ws.append(list(row))

    audit_ws = wb.create_sheet("_audit")
    audit_ws.sheet_state = "hidden"
    if audit_events:
        from ...domain.audit import audit_events_to_rows
        rows = audit_events_to_rows(audit_events)
        if rows:
            keys = list(rows[0].keys())
            for col_idx, name in enumerate(keys, start=1):
                cell = audit_ws.cell(row=1, column=col_idx, value=str(name))
                cell.font = cell.font.copy(bold=True)
            for row in rows:
                audit_ws.append([row.get(k, "") for k in keys])
        else:
            for col_idx, name in enumerate(["step_id", "operation"], start=1):
                cell = audit_ws.cell(row=1, column=col_idx, value=name)
                cell.font = cell.font.copy(bold=True)
    else:
        for col_idx, name in enumerate(["step_id", "operation"], start=1):
            cell = audit_ws.cell(row=1, column=col_idx, value=name)
            cell.font = cell.font.copy(bold=True)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
