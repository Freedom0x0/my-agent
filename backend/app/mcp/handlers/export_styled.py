"""tablex_export_styled handler."""
from __future__ import annotations

from datetime import datetime, timezone
from typing import Any

from openpyxl import load_workbook

from ...db import OutputRecord, init_db, insert_output
from ...domain.styled_export import StyleError, apply_styles
from ..schemas import ToolCall, ToolResult
from ._base import HandlerSpec, _db_path, _fail, _ok, _write_export_workbook, _output_id


def handle_export_styled(tool_call: ToolCall, session: Any) -> ToolResult:
    file_id = tool_call.input["file_id"]
    sheet = tool_call.input["sheet"]
    style = tool_call.input.get("style") or {}
    output_name = tool_call.input.get("output_name")
    if not output_name:
        return _fail("output_name 是必填参数")
    if not isinstance(style, dict):
        return _fail("style 必须是 dict")

    sheet_ref = f"{file_id}::{sheet}"
    if sheet_ref not in session.tables:
        return _fail(f"工作表未加载: {sheet_ref}")

    output_id = _output_id(session)
    output_path = session.output_dir / f"{output_id}.xlsx"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # 1) Write the standard workbook (all session tables + audit) first.
    _write_export_workbook(output_path, session.tables, session.audit_events)

    # 2) Reload, find the sheet that corresponds to sheet_ref, apply styles.
    wb = load_workbook(output_path)
    try:
        # The standard writer escapes "::" to "_of_".
        escaped_ref = sheet_ref.replace("::", "_of_")
        if escaped_ref in wb.sheetnames:
            target_title = escaped_ref
        elif sheet in wb.sheetnames:
            target_title = sheet
        else:
            return _fail(f"输出工作簿找不到 sheet: {sheet}")
        try:
            apply_styles(wb, target_title, style)
        except StyleError as exc:
            return _fail(str(exc))
        wb.save(output_path)
    finally:
        wb.close()

    session.output_id = output_id
    session.output_path = str(output_path)

    db_path = _db_path()
    init_db(db_path)
    insert_output(
        db_path,
        OutputRecord(
            output_id=output_id,
            stored_path=str(output_path),
            source_file_ids=[file_id],
            plan={"tool": "tablex_export_styled", "input": tool_call.input},
            result={"sheet_name": sheet},
            status="completed",
            created_at=datetime.now(timezone.utc).isoformat(),
            output_name=output_name,
        ),
    )

    return _ok(
        f"样式导出完成：{sheet}",
        # See export.py — the client needs the id to reopen this file.
        data={
            "output_id": output_id,
            "output_name": output_name,
            "sheets": [sheet],
        },
    )


HANDLERS = {
    "tablex_export_styled": HandlerSpec(
        "tablex_export_styled", ["file_id", "sheet", "style", "output_name"], handle_export_styled,
    ),
}
