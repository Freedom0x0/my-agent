"""Anti-corruption handlers: tool_use JSON → domain call → ToolResult.

Each handler signature is `def handle_xxx(tool_call, session) -> ToolResult`.
Handlers MUST NOT raise — they return `ToolResult(success=False, error=...)` on failure.
"""
from __future__ import annotations

import json
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook

from ..config import get_settings
from ..db import OutputRecord, init_db, insert_output
from ..domain.audit import make_audit_event
from ..domain.executor import (
    INTERNAL_COL,
    _apply_compare,
    _apply_deduplicate,
    _apply_fill_formula,
    _apply_fill_null,
    _apply_filter,
    _apply_group_summary,
    _apply_normalize,
    _apply_sort,
)
from ..domain._archive.operations import InvalidPlanError
from ..domain.parser import inspect_workbook, load_tables
from ..schemas import (
    CompareOperation,
    DeduplicateOperation,
    FillFormulaOperation,
    FilterCondition,
    FilterOperation,
    GroupSummaryOperation,
    NormalizeOperation,
)
from .schemas import MAX_RESULT_BYTES, ToolCall, ToolResult, ValidationResult


# ----- Registry -----

HandlerFn = Callable[[ToolCall, "Any"], ToolResult]


class HandlerSpec:
    """Declares what fields a tool requires and which function runs it."""

    def __init__(self, name: str, required: list[str], fn: HandlerFn):
        self.name = name
        self.required = required
        self.fn = fn


# Tools that take file_id + sheet (sheet must already be loaded).
_FILE_SHEET_TOOLS = {
    "tablex_normalize",
    "tablex_deduplicate",
    "tablex_filter",
    "tablex_group_summary",
    "tablex_fill_formula",
    "tablex_sort",
    "tablex_fill_null",
}

# Tools that take file_id only.
_FILE_TOOLS = {"tablex_inspect"}


def validate_tool_input(name: str, input_dict: dict, session: Any) -> ValidationResult:
    """Anti-corruption layer 1: parameter validation. Never enters the domain layer."""
    if name not in HANDLERS:
        return ValidationResult(ok=False, error=f"未知工具: {name}")

    spec = HANDLERS[name]
    for field in spec.required:
        if field not in input_dict or input_dict[field] in (None, ""):
            return ValidationResult(ok=False, error=f"缺少参数: {field}")

    if name in _FILE_SHEET_TOOLS or name in _FILE_TOOLS:
        file_id = input_dict.get("file_id")
        if not session.has_file(file_id):
            return ValidationResult(ok=False, error=f"文件不存在: {file_id}")

    if name in _FILE_SHEET_TOOLS:
        sheet = input_dict.get("sheet")
        if sheet is None:
            return ValidationResult(ok=False, error="缺少参数: sheet")
        sheet_ref = f"{file_id}::{sheet}"
        if sheet_ref not in session.tables:
            return ValidationResult(ok=False, error=f"工作表未加载: {sheet}")

    if name == "tablex_compare":
        for ref_field in ("left_ref", "right_ref"):
            ref = input_dict.get(ref_field)
            if ref not in session.tables:
                return ValidationResult(ok=False, error=f"工作表未加载: {ref}")
        for col in input_dict.get("key_columns") or []:
            left_ref = input_dict["left_ref"]
            if col not in session.tables[left_ref].columns:
                return ValidationResult(ok=False, error=f"列 {col} 不在 {left_ref}")
            right_ref = input_dict["right_ref"]
            if col not in session.tables[right_ref].columns:
                return ValidationResult(ok=False, error=f"列 {col} 不在 {right_ref}")

    return ValidationResult(ok=True)


def truncate_result(result: ToolResult, max_bytes: int = MAX_RESULT_BYTES) -> str:
    """Serialize ToolResult; if over budget, drop the data field (keep summary/error)."""
    text = result.model_dump_json(exclude_none=True)
    if len(text.encode("utf-8")) <= max_bytes:
        return text
    minimal = {k: v for k, v in {"success": result.success, "summary": result.summary, "error": result.error}.items() if v is not None}
    return json.dumps(minimal, ensure_ascii=False)


def _ok(summary: str, data: dict[str, Any] | None = None) -> ToolResult:
    return ToolResult(success=True, summary=summary, data=data)


def _fail(error: str, summary: str = "处理失败") -> ToolResult:
    return ToolResult(success=False, summary=summary, error=error[:300])


def _audit(
    session: Any,
    operation: str,
    sheet_ref: str,
    columns: list[str],
    input_rows: int,
    output_rows: int,
    affected: list[int],
    details: dict[str, Any] | None = None,
) -> None:
    event = make_audit_event(
        step_id=f"step-{len(session.audit_events) + 1:03d}",
        operation=operation,
        input_sheets=[sheet_ref],
        output_sheets=[sheet_ref],
        columns=columns,
        input_rows=input_rows,
        output_rows=output_rows,
        affected_rows=affected,
        details=details,
    )
    session.audit_events.append(event)


def _db_path() -> Path:
    base = Path(get_settings().APP_DATA_DIR).resolve()
    base.mkdir(parents=True, exist_ok=True)
    return base / "metadata.db"


# ----- Handlers -----


def handle_upload(tool_call: ToolCall, session: Any) -> ToolResult:
    file_id = tool_call.input["file_id"]
    info = session.files[file_id]
    path = Path(info["path"])
    if not path.exists():
        return _fail(f"文件已不存在: {path}")
    try:
        inspection = inspect_workbook(path)
        tables = load_tables(path)
    except Exception as exc:
        return _fail(f"无法解析文件: {exc}")

    loaded: list[str] = []
    for sheet_name, df in tables.items():
        ref = f"{file_id}::{sheet_name}"
        session.tables[ref] = df
        loaded.append(sheet_name)

    issue_count = sum(len(s.issues) for s in inspection.sheets)
    summary = f"已加载 {inspection.filename}，共 {len(inspection.sheets)} 个工作表，发现 {issue_count} 个数据质量问题"
    return _ok(
        summary,
        data={
            "file_id": file_id,
            "sheets": [s.name for s in inspection.sheets],
            "loaded_refs": [f"{file_id}::{n}" for n in loaded],
            "issue_count": issue_count,
        },
    )


def handle_inspect(tool_call: ToolCall, session: Any) -> ToolResult:
    file_id = tool_call.input["file_id"]
    sheet_name = tool_call.input.get("sheet")
    path = Path(session.files[file_id]["path"])
    try:
        inspection = inspect_workbook(path)
    except Exception as exc:
        return _fail(f"无法检查文件: {exc}")

    target = None
    if sheet_name:
        target = next((s for s in inspection.sheets if s.name == sheet_name), None)
        if target is None:
            return _fail(f"工作表不存在: {sheet_name}")
    elif inspection.sheets:
        target = inspection.sheets[0]

    if target is None:
        return _ok("文件没有任何工作表", data={"issues": []})

    issues = [
        {"code": i.code, "severity": i.severity, "message": i.message, "column": i.column}
        for i in target.issues
    ]
    summary = (
        f"工作表 '{target.name}': {target.row_count} 行 × {target.column_count} 列，"
        f"发现 {len(issues)} 个问题"
    )
    return _ok(
        summary,
        data={
            "sheet": target.name,
            "row_count": target.row_count,
            "column_count": target.column_count,
            "columns": [
                {"name": c.name, "type": c.inferred_type, "null_count": c.null_count}
                for c in target.columns
            ],
            "issues": issues[:20],
        },
    )


def handle_normalize(tool_call: ToolCall, session: Any) -> ToolResult:
    file_id = tool_call.input["file_id"]
    sheet = tool_call.input["sheet"]
    columns = tool_call.input["columns"]
    target_type = tool_call.input["target_type"]
    sheet_ref = f"{file_id}::{sheet}"
    df = session.tables[sheet_ref]

    missing = [c for c in columns if c not in df.columns]
    if missing:
        return _fail(f"列不存在: {missing}")

    op = NormalizeOperation(sheet=sheet_ref, columns=columns, target_type=target_type)
    new_df, affected = _apply_normalize(df, op)
    session.tables[sheet_ref] = new_df
    _audit(session, "normalize", sheet_ref, columns, len(df), len(new_df), affected, {"target_type": target_type})
    return _ok(
        f"已统一 {len(columns)} 列格式，影响 {len(affected)} 行",
        data={"columns": columns, "affected_rows": affected[:20]},
    )


def handle_deduplicate(tool_call: ToolCall, session: Any) -> ToolResult:
    file_id = tool_call.input["file_id"]
    sheet = tool_call.input["sheet"]
    key_columns = tool_call.input["key_columns"]
    keep = tool_call.input.get("keep", "first")
    sheet_ref = f"{file_id}::{sheet}"
    df = session.tables[sheet_ref]

    missing = [c for c in key_columns if c not in df.columns]
    if missing:
        return _fail(f"列不存在: {missing}")

    op = DeduplicateOperation(sheet=sheet_ref, key_columns=key_columns, keep=keep)
    new_df, affected = _apply_deduplicate(df, op)
    removed = len(df) - len(new_df)
    session.tables[sheet_ref] = new_df
    _audit(
        session, "deduplicate", sheet_ref, key_columns, len(df), len(new_df), affected,
        {"keep": keep, "removed": removed},
    )
    return _ok(
        f"删除 {removed} 行重复数据（基于 {', '.join(key_columns)}）",
        data={"removed_rows": affected[:20], "input_rows": len(df), "output_rows": len(new_df)},
    )


def handle_filter(tool_call: ToolCall, session: Any) -> ToolResult:
    file_id = tool_call.input["file_id"]
    sheet = tool_call.input["sheet"]
    conditions_raw = tool_call.input["conditions"]
    match = tool_call.input.get("match", "all")
    output_sheet = tool_call.input.get("output_sheet", "筛选结果")
    sheet_ref = f"{file_id}::{sheet}"
    df = session.tables[sheet_ref]

    conditions = []
    for c in conditions_raw:
        try:
            conditions.append(FilterCondition.model_validate(c))
        except Exception as exc:
            return _fail(f"筛选条件无效: {exc}")
        if conditions[-1].column not in df.columns:
            return _fail(f"列 {conditions[-1].column} 不存在")

    op = FilterOperation(
        sheet=sheet_ref, conditions=conditions, match=match, output_sheet=output_sheet,
    )
    filtered = _apply_filter(df, op)
    session.tables[output_sheet] = filtered
    _audit(
        session, "filter", sheet_ref,
        [c.column for c in conditions], len(df), len(filtered), [],
        {"conditions": [c.model_dump() for c in conditions], "match": match, "output_sheet": output_sheet},
    )
    return _ok(
        f"筛选出 {len(filtered)} 条记录，写入工作表 '{output_sheet}'",
        data={"output_sheet": output_sheet, "matched": len(filtered)},
    )


def handle_group_summary(tool_call: ToolCall, session: Any) -> ToolResult:
    file_id = tool_call.input["file_id"]
    sheet = tool_call.input["sheet"]
    group_by = tool_call.input["group_by"]
    metrics = tool_call.input["metrics"]
    output_sheet = tool_call.input.get("output_sheet", "汇总结果")
    sheet_ref = f"{file_id}::{sheet}"
    df = session.tables[sheet_ref]

    missing = [c for c in group_by if c not in df.columns]
    if missing:
        return _fail(f"分组列不存在: {missing}")
    for col in metrics.keys():
        if col not in df.columns:
            return _fail(f"聚合列不存在: {col}")

    op = GroupSummaryOperation(
        sheet=sheet_ref, group_by=group_by, metrics=metrics, output_sheet=output_sheet,
    )
    summary = _apply_group_summary(df, op)
    session.tables[output_sheet] = summary
    _audit(
        session, "group_summary", sheet_ref,
        list({*group_by, *metrics.keys()}), len(df), len(summary), [],
        {"group_by": group_by, "metrics": metrics, "output_sheet": output_sheet},
    )
    return _ok(
        f"已按 {', '.join(group_by)} 汇总为 {len(summary)} 组，写入 '{output_sheet}'",
        data={"output_sheet": output_sheet, "group_count": len(summary)},
    )


def handle_compare(tool_call: ToolCall, session: Any) -> ToolResult:
    left_ref = tool_call.input["left_ref"]
    right_ref = tool_call.input["right_ref"]
    key_columns = tool_call.input["key_columns"]
    output_sheet = tool_call.input.get("output_sheet", "对比结果")

    if left_ref not in session.tables:
        return _fail(f"工作表未加载: {left_ref}")
    if right_ref not in session.tables:
        return _fail(f"工作表未加载: {right_ref}")

    left = session.tables[left_ref]
    right = session.tables[right_ref]
    op = CompareOperation(
        left_sheet=left_ref, right_sheet=right_ref, key_columns=key_columns,
        compare_columns=None, output_sheet=output_sheet,
    )
    try:
        result_df = _apply_compare(left, right, op)
    except InvalidPlanError as exc:
        return _fail(str(exc))
    session.tables[output_sheet] = result_df

    counts = result_df["变化类型"].value_counts().to_dict() if "变化类型" in result_df.columns else {}
    added = int(counts.get("added", 0))
    removed = int(counts.get("removed", 0))
    changed = int(counts.get("changed", 0))
    _audit(
        session, "compare", f"{left_ref},{right_ref}", key_columns,
        len(left), len(result_df), [],
        {"left_ref": left_ref, "right_ref": right_ref, "key_columns": key_columns, "output_sheet": output_sheet},
    )
    return _ok(
        f"对比完成：新增 {added} 条、删除 {removed} 条、变化 {changed} 条",
        data={"output_sheet": output_sheet, "added": added, "removed": removed, "changed": changed},
    )


def handle_fill_formula(tool_call: ToolCall, session: Any) -> ToolResult:
    file_id = tool_call.input["file_id"]
    sheet = tool_call.input["sheet"]
    target_column = tool_call.input["target_column"]
    expression = tool_call.input["expression"]
    start_row = int(tool_call.input["start_row"])
    end_row = int(tool_call.input["end_row"])
    only_blank = bool(tool_call.input.get("only_blank", True))
    sheet_ref = f"{file_id}::{sheet}"
    df = session.tables[sheet_ref]

    if target_column not in df.columns:
        return _fail(f"目标列 {target_column} 不存在")
    if start_row > end_row:
        return _fail("起始行不能大于结束行")

    op = FillFormulaOperation(
        sheet=sheet_ref, target_column=target_column, expression=expression,
        start_row=start_row, end_row=end_row, only_blank=only_blank,
    )
    try:
        new_df, affected = _apply_fill_formula(df, op)
    except InvalidPlanError as exc:
        return _fail(str(exc))
    session.tables[sheet_ref] = new_df
    _audit(
        session, "fill_formula", sheet_ref, [target_column], len(df), len(new_df), affected,
        {"expression": expression, "range": [start_row, end_row], "only_blank": only_blank},
    )
    return _ok(
        f"已为 {target_column} 列补充 {len(affected)} 个公式",
        data={"target_column": target_column, "affected_rows": affected[:20]},
    )


def handle_sort(tool_call: ToolCall, session: Any) -> ToolResult:
    file_id = tool_call.input["file_id"]
    sheet = tool_call.input["sheet"]
    column = tool_call.input["column"]
    order = tool_call.input.get("order", "asc")
    sheet_ref = f"{file_id}::{sheet}"
    df = session.tables[sheet_ref]

    if column not in df.columns:
        return _fail(f"列 {column} 不存在")
    try:
        new_df, affected = _apply_sort(df, column=column, order=order)
    except InvalidPlanError as exc:
        return _fail(str(exc))
    session.tables[sheet_ref] = new_df
    _audit(
        session, "sort", sheet_ref, [column], len(df), len(new_df), affected,
        {"column": column, "order": order},
    )
    label = "降序" if order == "desc" else "升序"
    return _ok(f"已按 {column} {label}排序", data={"column": column, "order": order})


def handle_fill_null(tool_call: ToolCall, session: Any) -> ToolResult:
    file_id = tool_call.input["file_id"]
    sheet = tool_call.input["sheet"]
    column = tool_call.input["column"]
    method = tool_call.input["method"]
    value = tool_call.input.get("value")
    sheet_ref = f"{file_id}::{sheet}"
    df = session.tables[sheet_ref]

    if column not in df.columns:
        return _fail(f"列 {column} 不存在")
    try:
        new_df, affected = _apply_fill_null(df, column=column, method=method, value=value)
    except InvalidPlanError as exc:
        return _fail(str(exc))
    session.tables[sheet_ref] = new_df
    _audit(
        session, "fill_null", sheet_ref, [column], len(df), len(new_df), affected,
        {"column": column, "method": method, "value": value},
    )
    label = {"mean": "均值", "ffill": "前值", "bfill": "后值", "value": "指定值"}.get(method, method)
    return _ok(
        f"已用 {label} 填充 {column} 列的 {len(affected)} 个空值",
        data={"column": column, "method": method, "filled_count": len(affected)},
    )


def _write_export_workbook(
    output_path: Path,
    tables: dict,
    audit_events: list,
) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    taken: set[str] = set()

    def _escape(name: str) -> str:
        # Openpyxl forbids these chars in sheet titles: :\\/?*[] and 31-char limit.
        cleaned = (
            name.replace("::", "_of_")
            .replace(":", "_")
            .replace("\\", "_")
            .replace("/", "_")
            .replace("?", "_")
            .replace("*", "_")
            .replace("[", "_")
            .replace("]", "_")
            .strip()
        ) or "工作表"
        candidate = cleaned[:31]
        suffix = 1
        while candidate in taken:
            suffix += 1
            stem = cleaned[: max(0, 31 - len(str(suffix)) - 1)]
            candidate = f"{stem}_{suffix}"
        taken.add(candidate)
        return candidate

    for ref, df in tables.items():
        ws = wb.create_sheet(_escape(ref))
        cleaned = df.drop(columns=[INTERNAL_COL], errors="ignore") if INTERNAL_COL in df.columns else df
        for col_idx, name in enumerate(cleaned.columns, start=1):
            ws.cell(row=1, column=col_idx, value=str(name)).font = ws.cell(row=1, column=col_idx).font.copy(bold=True)
        for row in cleaned.itertuples(index=False, name=None):
            ws.append(list(row))

    audit_ws = wb.create_sheet("_audit")
    audit_ws.sheet_state = "hidden"
    if audit_events:
        from ..domain.audit import audit_events_to_rows
        rows = audit_events_to_rows(audit_events)
        if rows:
            keys = list(rows[0].keys())
            for col_idx, name in enumerate(keys, start=1):
                cell = audit_ws.cell(row=1, column=col_idx, value=str(name))
                cell.font = cell.font.copy(bold=True)
            for row in rows:
                audit_ws.append([row.get(k, "") for k in keys])
        else:
            for col_idx, name in enumerate(["step_id", "operation"], start=1):
                cell = audit_ws.cell(row=1, column=col_idx, value=name)
                cell.font = cell.font.copy(bold=True)
    else:
        for col_idx, name in enumerate(["step_id", "operation"], start=1):
            cell = audit_ws.cell(row=1, column=col_idx, value=name)
            cell.font = cell.font.copy(bold=True)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def handle_export(tool_call: ToolCall, session: Any) -> ToolResult:
    if not session.tables:
        return _fail("当前会话没有任何工作表可以导出")

    output_id = uuid.uuid4().hex
    output_path = session.output_dir / f"{output_id}.xlsx"
    _write_export_workbook(output_path, session.tables, session.audit_events)

    session.output_id = output_id
    session.output_path = str(output_path)

    db_path = _db_path()
    init_db(db_path)
    record = OutputRecord(
        output_id=output_id,
        stored_path=str(output_path),
        source_file_ids=list(session.files.keys()),
        plan={"tool_calls": session.tool_calls_log},
        result={
            "audit_events": [e.model_dump() for e in session.audit_events],
            "sheets": list(session.tables.keys()),
        },
        status="completed",
        created_at=datetime.now(timezone.utc).isoformat(),
    )
    insert_output(db_path, record)

    sheets = list(session.tables.keys())
    return _ok(
        f"已生成处理结果文件，共 {len(sheets)} 个工作表",
        data={"output_id": output_id, "sheets": sheets + ["_audit"]},
    )


HANDLERS: dict[str, HandlerSpec] = {
    "tablex_upload": HandlerSpec("tablex_upload", ["file_id"], handle_upload),
    "tablex_inspect": HandlerSpec("tablex_inspect", ["file_id"], handle_inspect),
    "tablex_normalize": HandlerSpec(
        "tablex_normalize", ["file_id", "sheet", "columns", "target_type"], handle_normalize,
    ),
    "tablex_deduplicate": HandlerSpec(
        "tablex_deduplicate", ["file_id", "sheet", "key_columns"], handle_deduplicate,
    ),
    "tablex_filter": HandlerSpec(
        "tablex_filter", ["file_id", "sheet", "conditions"], handle_filter,
    ),
    "tablex_group_summary": HandlerSpec(
        "tablex_group_summary",
        ["file_id", "sheet", "group_by", "metrics"], handle_group_summary,
    ),
    "tablex_compare": HandlerSpec(
        "tablex_compare", ["left_ref", "right_ref", "key_columns"], handle_compare,
    ),
    "tablex_fill_formula": HandlerSpec(
        "tablex_fill_formula",
        ["file_id", "sheet", "target_column", "expression", "start_row", "end_row"],
        handle_fill_formula,
    ),
    "tablex_sort": HandlerSpec(
        "tablex_sort", ["file_id", "sheet", "column"], handle_sort,
    ),
    "tablex_fill_null": HandlerSpec(
        "tablex_fill_null", ["file_id", "sheet", "column", "method"], handle_fill_null,
    ),
    "tablex_export": HandlerSpec("tablex_export", [], handle_export),
}