"""Chart business logic for tablex_chart."""
from __future__ import annotations

import io
from pathlib import Path
from typing import Any, Literal

import matplotlib
matplotlib.use("Agg")  # ponytail: headless backend, no DISPLAY needed.
import matplotlib.pyplot as plt
import pandas as pd

ChartType = Literal["bar", "line", "pie", "scatter", "heatmap"]


class ChartError(ValueError):
    """Raised when chart parameters are invalid."""


def auto_select_chart(df: pd.DataFrame, x: str, y: list[str]) -> str:
    if x not in df.columns:
        raise ChartError(f"x 列不存在: {x}")
    if not y:
        raise ChartError("y 列表不能为空")
    x_dtype_kind = _infer_kind(df[x])
    n_unique = df[x].dropna().nunique()
    n_y = len(y)

    if n_y == 1 and x_dtype_kind in {"text", "categorical"} and n_unique <= 8:
        return "pie"
    if x_dtype_kind == "datetime":
        return "line"
    if x_dtype_kind in {"text", "categorical"}:
        return "bar"
    if n_y >= 2 and x_dtype_kind == "number":
        return "scatter"
    if all(_infer_kind(df[c]) == "number" for c in y) and len(df) >= 20:
        return "heatmap"
    return "bar"


def _infer_kind(series: pd.Series) -> str:
    if pd.api.types.is_datetime64_any_dtype(series):
        return "datetime"
    if pd.api.types.is_bool_dtype(series):
        return "boolean"
    if pd.api.types.is_numeric_dtype(series):
        n_unique = series.dropna().nunique()
        return "number" if n_unique > 20 else "categorical"
    n_unique = series.dropna().nunique()
    return "categorical" if n_unique <= 20 else "text"


def generate_chart(
    df: pd.DataFrame,
    *,
    chart_type: str,
    x: str,
    y: list[str],
    title: str | None = None,
    style: dict[str, Any] | None = None,
) -> bytes:
    style = style or {}
    width = int(style.get("width", 800))
    height = int(style.get("height", 500))
    theme = style.get("theme", "light")

    if theme == "dark":
        plt.style.use("dark_background")
    else:
        plt.style.use("default")

    fig, ax = plt.subplots(figsize=(width / 100, height / 100))

    try:
        if chart_type == "bar":
            df.plot(kind="bar", x=x, y=y, ax=ax, title=title)
        elif chart_type == "line":
            df.plot(kind="line", x=x, y=y, ax=ax, title=title, marker="o")
        elif chart_type == "pie":
            if len(y) != 1:
                raise ChartError("饼图仅支持单个 y 列")
            data = df.groupby(x)[y[0]].sum().reset_index()
            ax.pie(data[y[0]], labels=data[x].astype(str).tolist(), autopct="%1.1f%%")
            if title:
                ax.set_title(title)
        elif chart_type == "scatter":
            if len(y) < 2:
                raise ChartError("散点图至少需要 2 个 y 列")
            df.plot(kind="scatter", x=y[0], y=y[1], ax=ax, title=title)
        elif chart_type == "heatmap":
            _draw_heatmap(df, x, y, ax, title)
        else:
            raise ChartError(f"不支持的图表类型: {chart_type}")

        if chart_type != "pie":
            ax.set_xlabel(x)
            ax.set_ylabel(", ".join(y))

        buf = io.BytesIO()
        fig.savefig(buf, format="png", bbox_inches="tight")
        return buf.getvalue()
    finally:
        plt.close(fig)


def _draw_heatmap(
    df: pd.DataFrame, x: str, y: list[str], ax, title: str | None,
) -> None:
    if len(y) < 1:
        raise ChartError("热力图至少需要 1 个 y 列")
    pivot = df.groupby(x)[y].mean()
    im = ax.imshow(pivot.values, aspect="auto", cmap="viridis")
    ax.set_xticks(range(len(pivot.columns)))
    ax.set_xticklabels([str(c) for c in pivot.columns], rotation=45, ha="right")
    ax.set_yticks(range(len(pivot.index)))
    ax.set_yticklabels([str(v) for v in pivot.index])
    ax.set_xlabel("")
    ax.set_ylabel(x)
    if title:
        ax.set_title(title)
    plt.colorbar(im, ax=ax)


def embed_in_excel(
    xlsx_path: str | Path,
    chart_png: bytes,
    sheet_name: str = "图表",
    cell: str = "A1",
) -> None:
    from openpyxl import load_workbook
    from openpyxl.drawing.image import Image as XLImage

    wb = load_workbook(str(xlsx_path))
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(sheet_name)
    else:
        ws = wb[sheet_name]
    img = XLImage(io.BytesIO(chart_png))
    img.anchor = cell
    ws.add_image(img)
    wb.save(str(xlsx_path))