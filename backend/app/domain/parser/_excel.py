"""File-format readers (xlsx / csv / xls) and public parser entry points."""
from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook as openpyxl_load

from ...schemas import ColumnInspection, DataIssue, SheetInspection, WorkbookInspection
from ._metadata import (
    INTERNAL_COL,
    INTERNAL_SOURCE_ROW,
    MAX_PREVIEW_ROWS,
    RESERVED_PREFIX,
    build_dataframe,
    detect_header_row,
    normalize_header,
)
from ._transform import infer_column_type
from ._validate import count_nulls, detect_duplicates, detect_mixed_date, detect_mixed_numeric

SUPPORTED_EXTENSIONS = {".xlsx", ".xls", ".csv"}

XLSX_SIGNATURE = b"PK"
XLS_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class UnsupportedFileError(ValueError):
    """Raised when the file extension or magic bytes are not supported."""


def _detect_extension(path: Path) -> str:
    return path.suffix.lower()


def _verify_magic(path: Path, ext: str) -> None:
    """Verify file header matches the declared extension."""
    raw = path.read_bytes()
    head = raw[:8]
    if ext == ".xlsx":
        if not head.startswith(XLSX_SIGNATURE):
            raise UnsupportedFileError("文件不是有效的 xlsx 格式")
    elif ext == ".xls":
        if not head.startswith(XLS_SIGNATURE):
            raise UnsupportedFileError("文件不是有效的 xls 格式")
    elif ext == ".csv":
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                raw.decode(encoding)
                return
            except UnicodeDecodeError:
                continue
        raise UnsupportedFileError("csv 文件编码无法识别，请使用 UTF-8 或 GB18030")


def _read_csv_bytes(path: Path) -> bytes:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            raw.decode(encoding)
            return raw
        except UnicodeDecodeError:
            continue
    raise UnsupportedFileError("csv 文件编码无法识别，请使用 UTF-8 或 GB18030")


def _sniff_delimiter(text: str) -> str:
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
        return dialect.delimiter
    except csv.Error:
        return ","


def _build_column_inspections(
    df: pd.DataFrame, columns: list[str], sheet_name: str, issues: list[DataIssue],
) -> list[ColumnInspection]:
    col_inspections: list[ColumnInspection] = []
    for col in columns:
        null_count = count_nulls(df[col]) if not df.empty else 0
        unique_count = int(df[col].dropna().astype(str).str.strip().nunique()) if not df.empty else 0
        inferred_type, _, _, samples = infer_column_type(df[col]) if not df.empty else ("empty", 0, 0, [])
        col_inspections.append(
            ColumnInspection(
                name=col,
                inferred_type=inferred_type,
                null_count=null_count,
                unique_count=unique_count,
                sample_values=samples,
            ),
        )
        if col.startswith(RESERVED_PREFIX):
            issues.append(
                DataIssue(
                    code="reserved_column",
                    severity="error",
                    message=f"列名 {col} 使用了系统保留前缀",
                    sheet=sheet_name,
                    column=col,
                ),
            )
        if null_count > 0:
            null_rows = [i + 1 for i, v in enumerate(df[col].tolist()) if v is None or str(v).strip() == ""]
            issues.append(
                DataIssue(
                    code="null_values",
                    severity="warning",
                    message=f"列 {col} 存在 {null_count} 个空值",
                    sheet=sheet_name,
                    column=col,
                    rows=null_rows[:200],
                ),
            )
        if inferred_type == "mixed":
            mixed, mixed_rows = detect_mixed_numeric(df[col])
            if mixed:
                issues.append(
                    DataIssue(
                        code="mixed_numeric",
                        severity="warning",
                        message=f"列 {col} 同时包含数字和文本",
                        sheet=sheet_name,
                        column=col,
                        rows=mixed_rows,
                    ),
                )
        # mixed date detection on text-looking date columns
        mixed_date, rows_md = detect_mixed_date(df[col]) if not df.empty else (False, [])
        if mixed_date:
            issues.append(
                DataIssue(
                    code="mixed_date",
                    severity="warning",
                    message=f"列 {col} 存在多种日期格式",
                    sheet=sheet_name,
                    column=col,
                    rows=rows_md,
                ),
            )
    return col_inspections


def _sheet_from_xlsx(path: Path, sheet_name: str) -> tuple[SheetInspection, pd.DataFrame]:
    wb = openpyxl_load(path, read_only=True, data_only=False)
    try:
        ws = wb[sheet_name]
        raw_rows: list[list[Any]] = []
        for row in ws.iter_rows(values_only=True):
            raw_rows.append(list(row))
    finally:
        wb.close()
    issues: list[DataIssue] = []
    header_idx = detect_header_row(raw_rows)
    if header_idx < 0:
        issues.append(
            DataIssue(
                code="empty_sheet",
                severity="warning",
                message="未在 20 行内找到表头",
                sheet=sheet_name,
            ),
        )
        return SheetInspection(
            name=sheet_name,
            row_count=0,
            column_count=0,
            columns=[],
            issues=issues,
        ), pd.DataFrame()

    used_names: set[str] = set()
    raw_header = raw_rows[header_idx]
    columns = [normalize_header(cell, used_names, issues, sheet_name) for cell in raw_header]

    # Drop fully empty rows below the header
    body = [r for r in raw_rows[header_idx + 1:] if any(c is not None and str(c).strip() != "" for c in r)]

    df = pd.DataFrame(body, columns=columns)
    df = df.astype(object).where(df.notna(), None)
    df.columns = columns

    # Drop columns that are completely empty
    keep = [c for c in columns if (df[c].notna() & (df[c].astype(str).str.strip() != "")).any()]
    df = df[keep]
    columns = keep

    formula_count = 0
    # detect formulas via second open of workbook
    wb2 = openpyxl_load(path, read_only=True, data_only=False)
    try:
        ws2 = wb2[sheet_name]
        for row in ws2.iter_rows(values_only=False):
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    formula_count += 1
    finally:
        wb2.close()

    col_inspections = _build_column_inspections(df, columns, sheet_name, issues)

    if not df.empty:
        dup, dup_rows = detect_duplicates(df)
        if dup:
            issues.append(
                DataIssue(
                    code="duplicate_rows",
                    severity="warning",
                    message=f"检测到 {len(dup_rows)} 行重复",
                    sheet=sheet_name,
                    rows=dup_rows,
                ),
            )

    preview_records = df.head(MAX_PREVIEW_ROWS).to_dict(orient="records") if not df.empty else []

    sheet = SheetInspection(
        name=sheet_name,
        row_count=len(df),
        column_count=len(columns),
        columns=col_inspections,
        preview=preview_records,
        issues=issues,
        formula_count=formula_count,
    )
    df_with_row = df.copy()
    df_with_row[INTERNAL_SOURCE_ROW] = list(range(header_idx + 2, header_idx + 2 + len(df)))
    return sheet, df_with_row


def _sheet_from_csv(path: Path) -> tuple[SheetInspection, pd.DataFrame]:
    raw = _read_csv_bytes(path)
    text = raw.decode("utf-8-sig", errors="replace")
    if text.startswith("﻿"):
        text = text[1:]
    delimiter = _sniff_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    raw_rows = [row for row in reader]
    issues: list[DataIssue] = []
    header_idx = detect_header_row(raw_rows)
    if header_idx < 0:
        issues.append(
            DataIssue(
                code="empty_sheet",
                severity="warning",
                message="未在 20 行内找到表头",
                sheet="数据",
            ),
        )
        return SheetInspection(
            name="数据",
            row_count=0,
            column_count=0,
            columns=[],
            issues=issues,
        ), pd.DataFrame()

    used: set[str] = set()
    columns = [normalize_header(cell, used, issues, "数据") for cell in raw_rows[header_idx]]
    body = [r for r in raw_rows[header_idx + 1:] if any(c is not None and str(c).strip() != "" for c in r)]
    df = pd.DataFrame(body, columns=columns)
    df = df.astype(object).where(df.notna(), None)
    keep = [c for c in columns if (df[c].notna() & (df[c].astype(str).str.strip() != "")).any()]
    df = df[keep]
    columns = keep

    col_inspections = _build_column_inspections(df, columns, "数据", issues)

    if not df.empty:
        dup, dup_rows = detect_duplicates(df)
        if dup:
            issues.append(
                DataIssue(
                    code="duplicate_rows",
                    severity="warning",
                    message=f"检测到 {len(dup_rows)} 行重复",
                    sheet="数据",
                    rows=dup_rows,
                ),
            )

    preview_records = df.head(MAX_PREVIEW_ROWS).to_dict(orient="records") if not df.empty else []
    sheet = SheetInspection(
        name="数据",
        row_count=len(df),
        column_count=len(columns),
        columns=col_inspections,
        preview=preview_records,
        issues=issues,
    )
    df_with_row = df.copy()
    df_with_row[INTERNAL_SOURCE_ROW] = list(range(header_idx + 2, header_idx + 2 + len(df)))
    return sheet, df_with_row


def _sheet_from_xls(path: Path) -> tuple[SheetInspection, pd.DataFrame]:
    xls = pd.ExcelFile(path, engine="xlrd")
    if not xls.sheet_names:
        return SheetInspection(
            name="Sheet1",
            row_count=0,
            column_count=0,
            columns=[],
            issues=[
                DataIssue(
                    code="empty_sheet",
                    severity="warning",
                    message="xls 工作簿没有可用工作表",
                    sheet="Sheet1",
                ),
            ],
        ), pd.DataFrame()
    sheet_name = xls.sheet_names[0]
    df = xls.parse(sheet_name=sheet_name, header=None)
    raw_rows: list[list[Any]] = df.astype(object).where(df.notna(), None).values.tolist()
    issues: list[DataIssue] = []
    header_idx = detect_header_row(raw_rows)
    if header_idx < 0:
        issues.append(
            DataIssue(
                code="empty_sheet",
                severity="warning",
                message="未在 20 行内找到表头",
                sheet=sheet_name,
            ),
        )
        return SheetInspection(
            name=sheet_name,
            row_count=0,
            column_count=0,
            columns=[],
            issues=issues,
        ), pd.DataFrame()
    used: set[str] = set()
    columns = [normalize_header(cell, used, issues, sheet_name) for cell in raw_rows[header_idx]]
    body = raw_rows[header_idx + 1:]
    df = pd.DataFrame(body, columns=columns).astype(object).where(
        pd.DataFrame(body, columns=columns).notna(), None,
    )
    keep = [c for c in columns if (df[c].notna() & (df[c].astype(str).str.strip() != "")).any()]
    df = df[keep]
    columns = keep
    col_inspections = _build_column_inspections(df, columns, sheet_name, issues)
    preview_records = df.head(MAX_PREVIEW_ROWS).to_dict(orient="records") if not df.empty else []
    sheet = SheetInspection(
        name=sheet_name,
        row_count=len(df),
        column_count=len(columns),
        columns=col_inspections,
        preview=preview_records,
        issues=issues,
        formula_count=0,
    )
    df_with_row = df.copy()
    df_with_row[INTERNAL_SOURCE_ROW] = list(range(header_idx + 2, header_idx + 2 + len(df)))
    return sheet, df_with_row


def inspect_workbook(path: Path) -> WorkbookInspection:
    if not path.exists():
        raise FileNotFoundError(f"文件不存在: {path}")
    ext = _detect_extension(path)
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileError(f"不支持的扩展名 {ext}")
    _verify_magic(path, ext)

    if ext == ".xlsx":
        wb = openpyxl_load(path, read_only=True, data_only=False)
        try:
            sheet_names = wb.sheetnames
        finally:
            wb.close()
        inspections: list[SheetInspection] = []
        for name in sheet_names:
            sheet, _ = _sheet_from_xlsx(path, name)
            inspections.append(sheet)
        return WorkbookInspection(filename=path.name, file_type="xlsx", sheets=inspections)
    if ext == ".csv":
        sheet, _ = _sheet_from_csv(path)
        return WorkbookInspection(filename=path.name, file_type="csv", sheets=[sheet])
    sheet, _ = _sheet_from_xls(path)
    return WorkbookInspection(filename=path.name, file_type="xls", sheets=[sheet])


def load_tables(path: Path) -> dict[str, pd.DataFrame]:
    """Load normalized tables with internal source-row column."""
    if not path.exists():
        raise FileNotFoundError(f"文件不存在: {path}")
    ext = _detect_extension(path)
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileError(f"不支持的扩展名 {ext}")
    _verify_magic(path, ext)

    if ext == ".xlsx":
        wb = openpyxl_load(path, read_only=True, data_only=False)
        try:
            sheet_names = wb.sheetnames
        finally:
            wb.close()
        tables: dict[str, pd.DataFrame] = {}
        for name in sheet_names:
            _, df = _sheet_from_xlsx(path, name)
            tables[name] = df.reset_index(drop=True)
        return tables
    if ext == ".csv":
        _, df = _sheet_from_csv(path)
        return {"数据": df.reset_index(drop=True)}
    _, df = _sheet_from_xls(path)
    return {df.attrs.get("sheet_name", "Sheet1"): df.reset_index(drop=True)}


def is_supported_file(path: Path) -> bool:
    return _detect_extension(path) in SUPPORTED_EXTENSIONS
