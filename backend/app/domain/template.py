"""Template filling: substitute {{key}} placeholders in a workbook.

ponytail: pure regex replace. Upgrade path: ship a Jinja-style renderer
with conditionals/loops when templates need branching logic.
"""
from __future__ import annotations

import re
from typing import Any

from openpyxl import Workbook

_PLACEHOLDER = re.compile(r"\{\{\s*(\w+)\s*\}\}")


class TemplateError(ValueError):
    """Raised when template parameters are invalid."""


def fill_template(
    workbook: Workbook,
    sheet_name: str,
    data: dict[str, Any],
) -> tuple[int, list[str]]:
    """Replace every {{key}} in the sheet with data[key] (stringified).

    Returns (replaced_count, missing_keys) where missing_keys lists keys
    referenced in the template but not present in `data`.
    """
    if sheet_name not in workbook.sheetnames:
        raise TemplateError(f"工作表不存在: {sheet_name}")

    ws = workbook[sheet_name]
    replaced = 0
    missing: set[str] = set()

    def repl(match: re.Match[str]) -> str:
        nonlocal replaced
        key = match.group(1)
        if key in data:
            replaced += 1
            return str(data[key])
        missing.add(key)
        return match.group(0)

    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and "{{" in v:
                cell.value = _PLACEHOLDER.sub(repl, v)
    return replaced, sorted(missing)
