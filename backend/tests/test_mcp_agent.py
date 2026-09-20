"""Tests for the agent loop (mocked MiniMax API).

The loop no longer executes tools: the model *compiles* a workflow graph via
`tablex_propose_workflow`, which gets stored and the turn stops for approval.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

from backend.app.db import get_workflow, init_db
from backend.app.mcp import agent as agent_module
from backend.app.mcp.agent import ChatError, process_chat
from backend.app.mcp.schemas import MAX_TOOL_CALLS
from backend.app.mcp.session import SessionStore
from backend.app.mcp.workflow import META_TOOL_NAME


class _FakeSettings:
    APP_DATA_DIR = "runtime"
    MODEL_BASE_URL = "https://example.test"
    MODEL_API_KEY = "fake-key"
    MODEL_NAME = "fake-model"


_GRAPH: dict[str, Any] = {
    "nodes": [
        {"id": "n_up", "label": "读取收支明细", "tool": "tablex_upload", "input": {"file_id": "file-001"}},
        {"id": "n_sum", "label": "按部门汇总", "tool": "tablex_group_summary", "input": {"group_by": ["部门"]}},
    ],
    "edges": [
        {"from_node": "n_up", "to_node": "n_sum", "to_param": "sheet"},
    ],
}


@pytest.fixture()
def sample_workbook(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "明细"
    ws.append(["部门", "订单号", "金额"])
    ws.append(["研发", "A001", "1,200"])
    ws.append(["销售", "A002", "800"])
    ws.append(["研发", "A003", "2,000"])
    out = tmp_path / "sample.xlsx"
    wb.save(out)
    return out


@pytest.fixture()
def session_store(tmp_path: Path) -> SessionStore:
    out = tmp_path / "outputs"
    out.mkdir(parents=True, exist_ok=True)
    return SessionStore(out)


@pytest.fixture()
def loaded_session(session_store: SessionStore, sample_workbook: Path):
    """A session with the fixture file registered and its tables pre-loaded."""
    sess = session_store.get_or_create("sid")
    sess.files["file-001"] = {
        "path": str(sample_workbook), "sha256": "x", "original_name": "sample.xlsx",
    }
    from backend.app.domain.parser import load_tables

    for name, df in load_tables(sample_workbook).items():
        sess.tables[f"file-001::{name}"] = df
    return sess


def _make_response(content: list[dict[str, Any]], stop_reason: str = "end_turn") -> dict[str, Any]:
    return {"stop_reason": stop_reason, "content": content}


def _text_block(text: str) -> dict[str, Any]:
    return {"type": "text", "text": text}


def _tool_use_block(tool_use_id: str, name: str, input_dict: dict[str, Any]) -> dict[str, Any]:
    return {"type": "tool_use", "id": tool_use_id, "name": name, "input": input_dict}


def test_process_chat_end_turn_no_tools(session_store: SessionStore, monkeypatch) -> None:
    monkeypatch.setattr(agent_module, "_default_chat_caller",
                        lambda **kwargs: _make_response([_text_block("Hello back")], stop_reason="end_turn"))
    response = process_chat("sid", "hi", [], store=session_store, settings=_FakeSettings())
    assert response.reply == "Hello back"
    assert response.error_code is None
    assert response.tool_calls == []
    # No graph yet, so the session stays in drafting.
    assert response.stage == "drafting"
    assert response.graph is None


# ----- the model still gets every tool definition -----


def test_process_chat_sends_node_type_definitions_plus_the_meta_tool(
    session_store: SessionStore, monkeypatch
) -> None:
    seen_tools: list[list[dict[str, Any]]] = []

    def capture(**kwargs):
        seen_tools.append(kwargs["tools"])
        return _make_response([_text_block("ok")], stop_reason="end_turn")

    monkeypatch.setattr(agent_module, "_default_chat_caller", capture)
    process_chat("sid", "hi", [], store=session_store, settings=_FakeSettings())

    names = [t["name"] for t in seen_tools[0]]
    assert META_TOOL_NAME in names
    # Every node type is still sent — as reference, not as actions.
    assert "tablex_upload" in names
    assert "tablex_group_summary" in names


# ----- nothing is executed -----


def test_process_chat_executes_nothing(
    session_store: SessionStore, loaded_session, monkeypatch
) -> None:
    """A directly-called tablex_* tool is handed back with an error result."""
    responses = [
        _make_response(
            [_tool_use_block("tu-1", "tablex_normalize", {
                "file_id": "file-001", "sheet": "明细", "columns": ["金额"], "target_type": "number",
            })],
            stop_reason="tool_use",
        ),
        _make_response([_text_block("重新规划")], stop_reason="end_turn"),
    ]
    seen_tool_results: list[str] = []

    def caller(**kwargs):
        last = kwargs["messages"][-1]
        blocks = last.get("content")
        if isinstance(blocks, list) and blocks and blocks[0].get("type") == "tool_result":
            seen_tool_results.append(blocks[0]["content"])
        return responses.pop(0)

    monkeypatch.setattr(agent_module, "_default_chat_caller", caller)

    response = process_chat("sid", "统一金额格式", ["file-001"], store=session_store, settings=_FakeSettings())
    assert response.reply == "重新规划"
    assert response.tool_calls[0]["tool"] == "tablex_normalize"
    assert response.tool_calls[0]["status"] == "error"
    assert "不能直接调用" in response.tool_calls[0]["summary"]

    # The handler never ran: no table mutation, no audit event.
    assert loaded_session.audit_events == []
    assert str(loaded_session.tables["file-001::明细"]["金额"].iloc[0]) == "1,200"

    # The model is told how to fix it.
    assert META_TOOL_NAME in seen_tool_results[0]


def test_process_chat_rejects_unknown_tool(session_store: SessionStore, monkeypatch) -> None:
    responses = [
        _make_response([_tool_use_block("tu-1", "tablex_nope", {})], stop_reason="tool_use"),
        _make_response([_text_block("重新规划")], stop_reason="end_turn"),
    ]
    monkeypatch.setattr(agent_module, "_default_chat_caller", lambda **kwargs: responses.pop(0))

    response = process_chat("sid", "test", [], store=session_store, settings=_FakeSettings())
    assert response.reply == "重新规划"
    assert response.tool_calls[0]["status"] == "error"
    assert response.tool_calls[0]["summary"] == "未知工具"


def test_process_chat_max_tool_calls_exceeded(session_store: SessionStore, monkeypatch) -> None:
    # Always answers with a direct (rejected) tool call → the loop must bail.
    monkeypatch.setattr(agent_module, "_default_chat_caller",
                        lambda **kwargs: _make_response(
                            [_tool_use_block("tu-1", "tablex_normalize", {})],
                            stop_reason="tool_use",
                        ))

    with pytest.raises(ChatError) as exc_info:
        process_chat("sid", "循环", [], store=session_store, settings=_FakeSettings())
    assert exc_info.value.code == "too_many_steps"
    assert len(session_store.get_or_create("sid").tool_calls_log) == MAX_TOOL_CALLS


# ----- submitting the graph -----


def test_process_chat_propose_workflow_stores_graph_and_stops(
    session_store: SessionStore, loaded_session, monkeypatch
) -> None:
    model_calls: list[list[dict[str, Any]]] = []

    def caller(**kwargs):
        model_calls.append(kwargs["messages"])
        return _make_response(
            [
                _text_block("我建议这样处理。"),
                _tool_use_block("tu-1", META_TOOL_NAME, _GRAPH),
            ],
            stop_reason="tool_use",
        )

    monkeypatch.setattr(agent_module, "_default_chat_caller", caller)

    response = process_chat("sid", "按部门汇总", ["file-001"], store=session_store, settings=_FakeSettings())

    assert len(model_calls) == 1, "提交图后必须停住，不再问模型"
    assert response.stage == "awaiting_approval"
    assert response.error_code is None
    assert loaded_session.stage == "awaiting_approval"
    assert [n["id"] for n in response.graph["nodes"]] == ["n_up", "n_sum"]
    assert loaded_session.tables["file-001::明细"] is not None  # untouched, not re-run
    assert loaded_session.audit_events == []
    assert response.tool_calls == [
        {"tool": META_TOOL_NAME, "status": "ok", "summary": "工作流已生成，等待用户批准。"}
    ]


def test_process_chat_invalid_graph_is_rejected_then_retried(
    session_store: SessionStore, monkeypatch
) -> None:
    bad = {"nodes": [{"id": "n1", "label": "汇总", "tool": "tablex_nope"}]}
    responses = [
        _make_response([_tool_use_block("tu-1", META_TOOL_NAME, bad)], stop_reason="tool_use"),
        _make_response([_tool_use_block("tu-2", META_TOOL_NAME, _GRAPH)], stop_reason="tool_use"),
    ]
    monkeypatch.setattr(agent_module, "_default_chat_caller", lambda **kwargs: responses.pop(0))

    response = process_chat("sid", "汇总", [], store=session_store, settings=_FakeSettings())
    assert response.tool_calls[0]["status"] == "error"
    assert response.tool_calls[0]["summary"] == "工作流图不合法"
    assert response.tool_calls[1]["status"] == "ok"
    assert response.stage == "awaiting_approval"


def test_process_chat_graph_is_injected_into_every_later_turn(
    session_store: SessionStore, monkeypatch
) -> None:
    seen_user_turns: list[str] = []
    responses = [
        _make_response([_tool_use_block("tu-1", META_TOOL_NAME, _GRAPH)], stop_reason="tool_use"),
        _make_response([_text_block("好的")], stop_reason="end_turn"),
    ]

    def caller(**kwargs):
        seen_user_turns.append(kwargs["messages"][-1]["content"])
        return responses.pop(0)

    monkeypatch.setattr(agent_module, "_default_chat_caller", caller)

    process_chat("sid", "按部门汇总", [], store=session_store, settings=_FakeSettings())
    process_chat("sid", "再导出一份", [], store=session_store, settings=_FakeSettings())

    assert "## 当前工作流" not in seen_user_turns[0]
    assert "## 当前工作流" in seen_user_turns[1]
    # 序号 + label + 状态 + 入边
    assert "1. [n_up] 读取收支明细" in seen_user_turns[1]
    assert "2. [n_sum] 按部门汇总" in seen_user_turns[1]
    assert "○" in seen_user_turns[1]
    assert "入边: 1 → sheet" in seen_user_turns[1]


def test_process_chat_history_stays_alternating_after_submission(
    session_store: SessionStore, monkeypatch
) -> None:
    """The graph turn must not leave two user messages in a row for the next call."""
    responses = [
        _make_response([_tool_use_block("tu-1", META_TOOL_NAME, _GRAPH)], stop_reason="tool_use"),
        _make_response([_text_block("好的")], stop_reason="end_turn"),
    ]
    monkeypatch.setattr(agent_module, "_default_chat_caller", lambda **kwargs: responses.pop(0))

    process_chat("sid", "图", [], store=session_store, settings=_FakeSettings())
    process_chat("sid", "再来", [], store=session_store, settings=_FakeSettings())

    roles = [m["role"] for m in session_store.get_or_create("sid").messages]
    assert roles == ["user", "assistant", "user", "assistant", "user", "assistant"]


# ----- the stage must not drift (PRD: 会话状态机落库，跨请求保持) -----


def _two_turn_setup(
    tmp_path: Path,
    sample_workbook: Path,
    monkeypatch,
    second_response: dict[str, Any],
) -> tuple[SessionStore, Path]:
    """Turn 1 submits the graph; turn 2 is `second_response` (never a graph).

    The last scripted response repeats, so a turn that loops forever still works.
    """
    db_path = tmp_path / "metadata.db"
    init_db(db_path)
    out = tmp_path / "outputs"
    out.mkdir(parents=True, exist_ok=True)
    store = SessionStore(out, db_path=db_path)
    session = store.get_or_create("sid")
    session.files["file-001"] = {
        "path": str(sample_workbook), "sha256": "x", "original_name": "sample.xlsx",
    }

    responses = [
        _make_response([_tool_use_block("tu-1", META_TOOL_NAME, _GRAPH)], stop_reason="tool_use"),
        second_response,
    ]

    def caller(**kwargs):
        return responses.pop(0) if len(responses) > 1 else responses[0]

    monkeypatch.setattr(agent_module, "_default_chat_caller", caller)

    first = process_chat("sid", "按部门汇总", ["file-001"], store=store, settings=_FakeSettings())
    assert first.stage == "awaiting_approval"
    assert get_workflow(db_path, "sid")["stage"] == "awaiting_approval"
    return store, db_path


def test_stage_stays_awaiting_approval_after_a_text_only_turn(
    tmp_path: Path, sample_workbook: Path, monkeypatch
) -> None:
    """Regression: turn 2 submits no graph — response, memory and DB must all agree."""
    store, db_path = _two_turn_setup(
        tmp_path, sample_workbook, monkeypatch,
        _make_response([_text_block("好的，还需要改哪里？")], stop_reason="end_turn"),
    )

    second = process_chat("sid", "这张图是什么意思？", [], store=store, settings=_FakeSettings())

    assert second.stage == "awaiting_approval"          # the response
    assert store.get_or_create("sid").stage == "awaiting_approval"  # memory
    assert get_workflow(db_path, "sid")["stage"] == "awaiting_approval"  # SQLite
    assert second.graph is not None  # the graph is still there to approve

    # A fresh store (eviction / restart) must not flip the stage back either.
    fresh = SessionStore(tmp_path / "outputs", db_path=db_path).get_or_create("sid")
    assert fresh.stage == "awaiting_approval"


def test_stage_stays_awaiting_approval_on_max_tokens(
    tmp_path: Path, sample_workbook: Path, monkeypatch
) -> None:
    store, db_path = _two_turn_setup(
        tmp_path, sample_workbook, monkeypatch,
        _make_response([_text_block("...")], stop_reason="max_tokens"),
    )

    second = process_chat("sid", "继续", [], store=store, settings=_FakeSettings())

    assert second.error_code == "model_truncated"
    assert second.stage == "awaiting_approval"
    assert get_workflow(db_path, "sid")["stage"] == "awaiting_approval"


def test_stage_stays_awaiting_approval_when_the_loop_gives_up(
    tmp_path: Path, sample_workbook: Path, monkeypatch
) -> None:
    store, db_path = _two_turn_setup(
        tmp_path, sample_workbook, monkeypatch,
        _make_response([_tool_use_block("tu-x", "tablex_normalize", {})], stop_reason="tool_use"),
    )

    with pytest.raises(ChatError) as exc_info:
        process_chat("sid", "循环", [], store=store, settings=_FakeSettings())

    assert exc_info.value.code == "too_many_steps"
    assert store.get_or_create("sid").stage == "awaiting_approval"
    assert get_workflow(db_path, "sid")["stage"] == "awaiting_approval"


def test_no_graph_keeps_drafting_and_stores_nothing(tmp_path: Path, monkeypatch) -> None:
    db_path = tmp_path / "metadata.db"
    init_db(db_path)
    out = tmp_path / "outputs"
    out.mkdir(parents=True, exist_ok=True)
    store = SessionStore(out, db_path=db_path)
    monkeypatch.setattr(agent_module, "_default_chat_caller",
                        lambda **kwargs: _make_response([_text_block("ok")], stop_reason="end_turn"))

    response = process_chat("sid", "hi", [], store=store, settings=_FakeSettings())

    assert response.stage == "drafting"
    assert get_workflow(db_path, "sid") is None


# ----- setup / error paths -----


def test_process_chat_model_not_configured(session_store: SessionStore) -> None:
    class BadSettings:
        APP_DATA_DIR = "runtime"
        MODEL_BASE_URL = ""
        MODEL_API_KEY = ""
        MODEL_NAME = ""

    with pytest.raises(ChatError) as exc_info:
        process_chat("sid", "hi", [], store=session_store, settings=BadSettings())
    assert exc_info.value.code == "model_not_configured"


def test_process_chat_max_tokens_stop_reason(session_store: SessionStore, monkeypatch) -> None:
    monkeypatch.setattr(agent_module, "_default_chat_caller",
                        lambda **kwargs: _make_response([_text_block("...")], stop_reason="max_tokens"))
    response = process_chat("sid", "hi", [], store=session_store, settings=_FakeSettings())
    assert response.error_code == "model_truncated"


def test_process_chat_unknown_stop_reason_raises(session_store: SessionStore, monkeypatch) -> None:
    monkeypatch.setattr(agent_module, "_default_chat_caller",
                        lambda **kwargs: _make_response([], stop_reason="wat"))
    with pytest.raises(ChatError) as exc_info:
        process_chat("sid", "hi", [], store=session_store, settings=_FakeSettings())
    assert exc_info.value.code == "model_unknown_stop_reason"


def test_process_chat_empty_message_raises(session_store: SessionStore, monkeypatch) -> None:
    monkeypatch.setattr(agent_module, "_default_chat_caller",
                        lambda **kwargs: _make_response([_text_block("ok")], stop_reason="end_turn"))
    with pytest.raises(ChatError) as exc_info:
        process_chat("sid", "   ", [], store=session_store, settings=_FakeSettings())
    assert exc_info.value.code == "invalid_request"


def test_process_chat_file_not_found_raises(session_store: SessionStore, monkeypatch) -> None:
    monkeypatch.setattr(agent_module, "_default_chat_caller",
                        lambda **kwargs: _make_response([_text_block("ok")], stop_reason="end_turn"))
    with pytest.raises(ChatError) as exc_info:
        process_chat("sid", "hi", ["missing-file"], store=session_store, settings=_FakeSettings())
    assert exc_info.value.code == "file_not_found"


def test_process_chat_first_message_includes_file_context(
    session_store: SessionStore, sample_workbook: Path, monkeypatch
) -> None:
    sess = session_store.get_or_create("sid")
    sess.files["file-001"] = {
        "path": str(sample_workbook), "sha256": "x", "original_name": "sample.xlsx",
    }

    seen_messages: list[list[dict[str, Any]]] = []

    def capture(**kwargs):
        seen_messages.append(kwargs["messages"])
        return _make_response([_text_block("ok")], stop_reason="end_turn")

    monkeypatch.setattr(agent_module, "_default_chat_caller", capture)

    process_chat("sid", "汇总金额", ["file-001"], store=session_store, settings=_FakeSettings())
    first_user = seen_messages[0][0]
    assert "file-001" in first_user["content"]
    assert "汇总金额" in first_user["content"]


def test_set_chat_caller_restores_default(monkeypatch) -> None:
    from backend.app.mcp.agent import set_chat_caller

    custom = lambda **kwargs: {"stop_reason": "end_turn", "content": []}
    set_chat_caller(custom)
    assert agent_module._default_chat_caller is custom
    set_chat_caller(None)
    assert agent_module._default_chat_caller is agent_module.minimax_chat