from pathlib import Path

import pytest
from openpyxl import Workbook

from backend.app.domain.parser import (
    INTERNAL_SOURCE_ROW,
    UnsupportedFileError,
    inspect_workbook,
    load_tables,
)


def test_inspection_detects_duplicate_rows_and_nulls(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "人员"
    sheet.append(["部门", "姓名", "金额"])
    sheet.append(["研发", "张三", "1,200"])
    sheet.append(["研发", "张三", "1,200"])
    sheet.append(["", "李四", None])
    source = tmp_path / "messy.xlsx"
    workbook.save(source)

    result = inspect_workbook(source)

    assert result.sheets[0].row_count == 3
    codes = {issue.code for issue in result.sheets[0].issues}
    assert "duplicate_rows" in codes
    assert "null_values" in codes


def test_inspection_csv_supports_utf8_and_gb18030(tmp_path: Path) -> None:
    csv_path = tmp_path / "utf.csv"
    csv_path.write_text("部门,姓名,金额\n研发,张三,1200\n销售,李四,800\n", encoding="utf-8-sig")
    result = inspect_workbook(csv_path)
    assert result.file_type == "csv"
    assert result.sheets[0].row_count == 2

    gb_path = tmp_path / "gb.csv"
    gb_path.write_text("部门,姓名,金额\n研发,张三,1200\n", encoding="gb18030")
    result2 = inspect_workbook(gb_path)
    assert result2.sheets[0].row_count == 1


def test_inspection_handles_empty_sheet(tmp_path: Path) -> None:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "空"
    source = tmp_path / "empty.xlsx"
    workbook.save(source)

    result = inspect_workbook(source)
    sheet = result.sheets[0]
    assert sheet.row_count == 0
    assert any(issue.code == "empty_sheet" for issue in sheet.issues)


def test_inspection_missing_file(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        inspect_workbook(tmp_path / "absent.xlsx")


def test_inspection_unsupported_extension(tmp_path: Path) -> None:
    bad = tmp_path / "note.txt"
    bad.write_text("hello", encoding="utf-8")
    with pytest.raises(UnsupportedFileError) as exc:
        inspect_workbook(bad)
    assert ".txt" in str(exc.value)


def test_inspection_rejects_bad_xlsx_signature(tmp_path: Path) -> None:
    bad = tmp_path / "fake.xlsx"
    bad.write_bytes(b"NOT_A_ZIP_FILE")
    with pytest.raises(UnsupportedFileError):
        inspect_workbook(bad)


def test_inspection_detects_mixed_numeric_and_reserved_columns(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "明细"
    sheet.append(["部门", "金额"])
    sheet.append(["研发", "1,200"])
    sheet.append(["研发", "N/A"])
    sheet.append(["研发", "3000"])
    source = tmp_path / "mixed.xlsx"
    workbook.save(source)

    result = inspect_workbook(source)
    codes = {issue.code for issue in result.sheets[0].issues}
    assert "mixed_numeric" in codes


def test_load_tables_attaches_source_row(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "明细"
    sheet.append(["部门", "金额"])
    sheet.append(["研发", "100"])
    sheet.append(["销售", "200"])
    source = tmp_path / "data.xlsx"
    workbook.save(source)

    tables = load_tables(source)
    df = tables["明细"]
    assert INTERNAL_SOURCE_ROW in df.columns
    assert df[INTERNAL_SOURCE_ROW].tolist() == [2, 3]