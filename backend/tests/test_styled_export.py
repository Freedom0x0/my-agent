"""Tests for tablex_export_styled handler + styled_export domain."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from backend.app.domain.styled_export import StyleError, apply_styles
from backend.app.mcp.handlers import handle_export_styled
from backend.app.mcp.schemas import ToolCall
from backend.app.mcp.session import Session


@pytest.fixture()
def session(tmp_path: Path) -> Session:
    s = Session("s", tmp_path / "out")
    s.files["file-S"] = {"path": "x", "sha256": "x", "original_name": "s.xlsx"}
    s.tables["file-S::data"] = pd.DataFrame({
        "name": ["Alice", "Bob", "Charlie"],
        "amount": [500, 1500, 2000],
        "category": ["A", "B", "A"],
    })
    return s


# ---------- domain ----------


def test_apply_styles_header_bold(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["name", "amount"])
    ws.append(["a", 1])
    apply_styles(wb, "S", {"header": {"bold": True, "bg_color": "2b4a8b", "font_color": "ffffff"}})
    cell = ws.cell(row=1, column=1)
    assert cell.font.bold is True
    assert cell.fill.fgColor.rgb is not None


def test_apply_styles_freeze(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["h"])
    ws.append(["a"])
    apply_styles(wb, "S", {"freeze_header": True})
    assert ws.freeze_panes == "A2"


def test_apply_styles_auto_width(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["short", "this is a longer header value"])
    ws.append(["x", "y"])
    apply_styles(wb, "S", {"auto_column_width": True})
    assert ws.column_dimensions["B"].width > 10


def test_apply_styles_merge(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["h1", "h2"])
    apply_styles(wb, "S", {"merge_cells": [{"range": "A1:B1", "value": "title"}]})
    assert "A1:B1" in ws.merged_cells
    assert ws["A1"].value == "title"


def test_apply_styles_unknown_sheet(tmp_path: Path) -> None:
    wb = Workbook()
    wb.active.title = "S"
    with pytest.raises(StyleError):
        apply_styles(wb, "Nope", {"header": {"bold": True}})


# ---------- handler ----------


def test_handle_export_styled_basic(session: Session) -> None:
    res = handle_export_styled(
        ToolCall(
            tool_use_id="s1", name="tablex_export_styled",
            input={
                "file_id": "file-S", "sheet": "data",
                "style": {
                    "header": {"bold": True, "bg_color": "2b4a8b", "font_color": "ffffff"},
                    "highlight_rules": [
                        {"column": "amount", "condition": ">1000", "bg_color": "fff3cd"},
                    ],
                    "freeze_header": True,
                    "auto_column_width": True,
                },
            },
        ),
        session,
    )
    assert res.success
    assert res.data["output_id"]
    out_path = session.output_dir / f"{res.data['output_id']}.xlsx"
    assert out_path.exists()
    # Reopen and confirm styles survived. The writer escapes "::" to "_of_".
    wb2 = load_workbook(out_path)
    styled = wb2["file-S_of_data"]
    assert styled.freeze_panes == "A2"
    # Header bold applied.
    assert styled.cell(row=1, column=1).font.bold is True


def test_handle_export_styled_invalid_style_type(session: Session) -> None:
    res = handle_export_styled(
        ToolCall(
            tool_use_id="s1", name="tablex_export_styled",
            input={"file_id": "file-S", "sheet": "data", "style": "bad"},
        ),
        session,
    )
    assert not res.success


def test_handle_export_styled_unknown_sheet(session: Session) -> None:
    res = handle_export_styled(
        ToolCall(
            tool_use_id="s1", name="tablex_export_styled",
            input={"file_id": "file-S", "sheet": "nope", "style": {}},
        ),
        session,
    )
    assert not res.success
