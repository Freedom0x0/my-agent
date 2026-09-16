"""Tests for tablex_join handler and join_tables domain."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from backend.app.mcp.handlers import handle_join
from backend.app.mcp.schemas import ToolCall
from backend.app.mcp.session import Session, SessionStore
from backend.app.domain.joiner import JoinError, join_tables


@pytest.fixture()
def session(tmp_path: Path) -> Session:
    store = SessionStore(tmp_path / "out")
    store.output_dir.mkdir(parents=True, exist_ok=True)
    s = store.get_or_create("sess-join")

    wb1 = Workbook()
    ws = wb1.active
    ws.title = "users"
    ws.append(["id", "name"])
    ws.append([1, "Alice"])
    ws.append([2, "Bob"])
    ws.append([3, "Carol"])
    p1 = tmp_path / "users.xlsx"
    wb1.save(p1)
    s.files["file-L"] = {"path": str(p1), "sha256": "x", "original_name": "users.xlsx"}

    wb2 = Workbook()
    ws = wb2.active
    ws.title = "orders"
    ws.append(["uid", "amount"])
    ws.append([1, 100])
    ws.append([2, 200])
    ws.append([4, 400])
    p2 = tmp_path / "orders.xlsx"
    wb2.save(p2)
    s.files["file-R"] = {"path": str(p2), "sha256": "y", "original_name": "orders.xlsx"}

    s.tables["file-L::users"] = pd.DataFrame({"id": [1, 2, 3], "name": ["Alice", "Bob", "Carol"]})
    s.tables["file-R::orders"] = pd.DataFrame({"uid": [1, 2, 4], "amount": [100, 200, 400]})
    return s


# ---------- domain.join_tables ----------


def test_join_tables_inner() -> None:
    left = pd.DataFrame({"id": [1, 2, 3], "name": ["a", "b", "c"]})
    right = pd.DataFrame({"id": [1, 2, 4], "v": [10, 20, 40]})
    merged, meta = join_tables(left, right, on="id", how="inner")
    assert len(merged) == 2
    assert meta.matched == 2
    assert meta.unmatched_left == 1
    assert meta.unmatched_right == 1


def test_join_tables_left() -> None:
    left = pd.DataFrame({"id": [1, 2, 3], "name": ["a", "b", "c"]})
    right = pd.DataFrame({"id": [1, 2], "v": [10, 20]})
    merged, meta = join_tables(left, right, on="id", how="left")
    assert len(merged) == 3
    assert meta.matched == 2


def test_join_tables_left_on_right_on() -> None:
    left = pd.DataFrame({"id": [1, 2], "name": ["a", "b"]})
    right = pd.DataFrame({"uid": [1, 2], "v": [10, 20]})
    merged, meta = join_tables(left, right, left_on="id", right_on="uid", how="inner")
    assert len(merged) == 2


def test_join_tables_self_join() -> None:
    df = pd.DataFrame({"id": [1, 2], "ref": [2, None]})
    left = df.rename(columns={"id": "lid"})
    right = df.rename(columns={"id": "rid"})
    merged, meta = join_tables(left, right, left_on="ref", right_on="rid", how="left")
    assert len(merged) == 2


def test_join_tables_invalid_params() -> None:
    left = pd.DataFrame({"id": [1]})
    right = pd.DataFrame({"id": [1]})
    with pytest.raises(JoinError):
        join_tables(left, right)


def test_join_tables_missing_column() -> None:
    left = pd.DataFrame({"x": [1]})
    right = pd.DataFrame({"id": [1]})
    with pytest.raises(JoinError):
        join_tables(left, right, on="id")


# ---------- handler ----------


def test_handle_join_inner(session: Session) -> None:
    res = handle_join(
        ToolCall(
            tool_use_id="j1", name="tablex_join",
            input={
                "left": {"file_id": "file-L", "sheet": "users"},
                "right": {"file_id": "file-R", "sheet": "orders"},
                "left_on": "id", "right_on": "uid",
                "how": "inner",
            },
        ),
        session,
    )
    assert res.success
    assert res.data is not None
    assert res.data["rows"] == 2
    assert "join结果" in session.tables


def test_handle_join_left(session: Session) -> None:
    res = handle_join(
        ToolCall(
            tool_use_id="j1", name="tablex_join",
            input={
                "left": {"file_id": "file-L", "sheet": "users"},
                "right": {"file_id": "file-R", "sheet": "orders"},
                "left_on": "id", "right_on": "uid",
                "how": "left",
            },
        ),
        session,
    )
    assert res.success
    assert res.data["rows"] == 3


def test_handle_join_missing_sheet(session: Session) -> None:
    res = handle_join(
        ToolCall(
            tool_use_id="j1", name="tablex_join",
            input={
                "left": {"file_id": "file-L", "sheet": "nope"},
                "right": {"file_id": "file-R", "sheet": "orders"},
                "on": "id",
            },
        ),
        session,
    )
    assert not res.success


def test_handle_join_no_key(session: Session) -> None:
    res = handle_join(
        ToolCall(
            tool_use_id="j1", name="tablex_join",
            input={
                "left": {"file_id": "file-L", "sheet": "users"},
                "right": {"file_id": "file-R", "sheet": "orders"},
            },
        ),
        session,
    )
    assert not res.success


def test_handle_join_writes_output_file(session: Session) -> None:
    res = handle_join(
        ToolCall(
            tool_use_id="j1", name="tablex_join",
            input={
                "left": {"file_id": "file-L", "sheet": "users"},
                "right": {"file_id": "file-R", "sheet": "orders"},
                "left_on": "id", "right_on": "uid",
                "how": "inner",
            },
        ),
        session,
    )
    assert res.success
    assert res.data is not None
    out_id = res.data["output_id"]
    assert (session.output_dir / f"{out_id}.xlsx").exists()