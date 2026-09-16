"""Tests for tablex_formula_graph handler + formula domain."""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from backend.app.domain.formula import (
    FormulaInfo,
    build_dependency_graph,
    extract_formulas,
    find_circular,
)
from backend.app.mcp.handlers import handle_formula_graph
from backend.app.mcp.schemas import ToolCall
from backend.app.mcp.session import Session


def _make_workbook(tmp_path: Path, with_circular: bool = False) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 10
    ws["B1"] = 20
    ws["C1"] = "=A1+B1"
    ws["A2"] = "=C1*2"
    ws["B2"] = "=A2+1"
    if with_circular:
        # C1 already depends on A1+B1. Add A1 = =C1 -> A1 <-> C1 cycle.
        ws["A1"] = "=C1+1"
    out = tmp_path / "f.xlsx"
    wb.save(out)
    return out


# ---------- domain ----------


def test_extract_formulas(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path)
    wb = Workbook()  # placeholder, will reload below
    from openpyxl import load_workbook
    wb = load_workbook(path)
    out = extract_formulas(wb, "Sheet1")
    cells = {f.cell for f in out}
    assert {"C1", "A2", "B2"}.issubset(cells)


def test_build_dependency_graph() -> None:
    formulas = [
        FormulaInfo("C1", "=A1+B1"),
        FormulaInfo("A2", "=C1*2"),
    ]
    g = build_dependency_graph(formulas)
    assert g["C1"] == ["A1", "B1"]
    assert g["A2"] == ["C1"]


def test_find_circular(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path, with_circular=True)
    from openpyxl import load_workbook
    wb = load_workbook(path)
    formulas = extract_formulas(wb, "Sheet1")
    graph = build_dependency_graph(formulas)
    cycles = find_circular(graph)
    assert any("A1" in c and "C1" in c for c in cycles)


def test_find_no_circular(tmp_path: Path) -> None:
    path = _make_workbook(tmp_path)
    from openpyxl import load_workbook
    wb = load_workbook(path)
    formulas = extract_formulas(wb, "Sheet1")
    graph = build_dependency_graph(formulas)
    cycles = find_circular(graph)
    assert cycles == []


# ---------- handler ----------


@pytest.fixture()
def session(tmp_path: Path) -> Session:
    s = Session("s", tmp_path / "out")
    path = _make_workbook(tmp_path)
    s.files["file-F"] = {"path": str(path), "sha256": "x", "original_name": "f.xlsx"}
    return s


def test_handle_formula_graph_full(session: Session) -> None:
    res = handle_formula_graph(
        ToolCall(
            tool_use_id="g1", name="tablex_formula_graph",
            input={"file_id": "file-F", "sheet": "Sheet1"},
        ),
        session,
    )
    assert res.success
    assert res.data["formula_count"] >= 3
    assert "C1" in res.data["graph"]


def test_handle_formula_graph_for_cell(session: Session) -> None:
    res = handle_formula_graph(
        ToolCall(
            tool_use_id="g1", name="tablex_formula_graph",
            input={"file_id": "file-F", "sheet": "Sheet1", "cell": "A2"},
        ),
        session,
    )
    assert res.success
    assert "C1" in res.data["upstream"]
    assert "B2" in res.data["downstream"]


def test_handle_formula_graph_unknown_file(session: Session) -> None:
    res = handle_formula_graph(
        ToolCall(
            tool_use_id="g1", name="tablex_formula_graph",
            input={"file_id": "missing", "sheet": "Sheet1"},
        ),
        session,
    )
    assert not res.success
