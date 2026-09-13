"""Create reproducible messy demo workbooks for the agent.

Run: python scripts/create_fixtures.py
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def _style_header(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="left")
        cell.fill = PatternFill("solid", fgColor="DCE6F1")


def create_personnel_messy(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "人员信息"
    ws.append(["姓名", "部门", "联系方式", "证件号", "入职日期"])
    rows = [
        ["张三", "研发部", "138-0000-0001", "110101199003078811", "2020/3/15"],
        ["张三", "研发部", "13800000001", "110101199003078811", "2020-03-15"],  # duplicate
        ["李四", "销售部", "139-0000-0002", "110101199205123424", "2019年7月1日"],
        ["王五", "研发部", None, "110101198811223344", ""],
        ["赵六", "", "137-0000-0003", "", "2021.10.01"],
        ["钱七", "销售部", "136-0000-0004", "110101199503156621", None],
        ["孙八", "研发部", "135-0000-0005", "110101198712091137", "2018-12-09"],
        ["孙八", "研发部", "13500000005", "110101198712091137", "2018/12/9"],  # duplicate
    ]
    for row in rows:
        ws.append(row)
    _style_header(ws)
    wb.save(path)


def create_project_progress_messy(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "项目进度"
    ws.append(["项目编号", "项目名称", "负责人", "状态", "截止日期"])
    rows = [
        ["P001", "客户系统迁移", "张三", "已完成", "2026-03-31"],
        ["P002", "数据平台升级", "李四", "进行中", "2026/06/15"],
        ["P003", "AI 应用集成", "王五", "延期", "2025年12月30日"],
        ["P004", "权限模块重构", "赵六", "未启动", "2026.09.01"],
        ["P005", "报表自动化", "钱七", "进行中", "2026-02-28"],
        ["P006", "数据治理", "孙八", "延期", "2025-11-30"],
        ["P007", "前端升级", "周九", "已完成", "2026-01-31"],
        ["P008", "移动端适配", "吴十", "进行中", "2026.07.31"],
    ]
    for row in rows:
        ws.append(row)
    _style_header(ws)
    wb.save(path)


def create_expenses_and_budget(path: Path) -> None:
    wb = Workbook()
    detail = wb.active
    detail.title = "收支明细"
    detail.append(["部门", "项目", "日期", "金额", "类别"])
    rows = [
        ["研发", "P001", "2026-01-05", "1,200.50", "差旅"],
        ["研发", "P001", "2026-01-08", "850.00", "差旅"],
        ["销售", "P002", "2026/02/12", "5,600", "市场"],
        ["销售", "P003", "2026年2月20日", "3,000元", "市场"],
        ["行政", "P004", "2026-03-02", "1,250", "办公"],
        ["行政", "P005", "2026.03.18", "450", "办公"],
        ["研发", "P002", "2026-04-09", "9,200", "采购"],
        ["销售", "P006", "2026-04-21", "2,800.00", "市场"],
        ["行政", "P007", "2026/05/14", "1,500", "办公"],
        ["研发", "P003", "2026.05.30", "6,300元", "采购"],
    ]
    for row in rows:
        detail.append(row)
    _style_header(detail)

    budget = wb.create_sheet("预算")
    budget.append(["部门", "预算金额"])
    for r in [("研发", "18000"), ("销售", "12000"), ("行政", "3200")]:
        budget.append(list(r))
    _style_header(budget)

    wb.save(path)


def main() -> None:
    fixtures_dir = Path(__file__).resolve().parent.parent / "fixtures"
    fixtures_dir.mkdir(parents=True, exist_ok=True)
    create_personnel_messy(fixtures_dir / "personnel_messy.xlsx")
    create_project_progress_messy(fixtures_dir / "project_progress_messy.xlsx")
    create_expenses_and_budget(fixtures_dir / "expenses_and_budget.xlsx")
    print(f"fixtures written to {fixtures_dir}")


if __name__ == "__main__":
    main()